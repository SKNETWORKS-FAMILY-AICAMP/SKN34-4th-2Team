"""Offline regression checks for regulation-style policy uploads."""
from io import BytesIO
import unittest

from docx import Document

from chatbot.cohort_document_rag import document_records
from vectordb.policy_ingestion import build_records, load_docx, infer_cohort


def regulation_bytes(articles=None, title="SKN34 훈련생 운영규정(안)"):
    doc = Document()
    doc.add_paragraph(title, style="Title")
    doc.add_heading("제1장 총칙", level=1)
    for heading, paragraphs in articles or [
        ("제1조(적용 범위)", ["이 규정은 해당 기수에 적용한다."]),
        ("제2조(출결)", ["① 다음 각 호는 지급 제외 사유이다.",
                       "1. 주 15시간 이상 근로하는 경우", "2. 그 밖에 정한 경우",
                       "② 다만, 제3조에 따른 예외를 적용한다."]),
        ("제3조(예외)", ["증빙을 제출하고 승인을 받은 경우에 한한다."]),
    ]:
        doc.add_heading(heading, level=2)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def upload_records(data, cohort="cohort_34", upload="a"):
    return document_records(data, "rules.docx", cohort, "policy",
                            f"cohorts/{cohort}/policy/rag/{upload * 32}.docx")


class RegulationIngestionTests(unittest.TestCase):
    def test_article_keeps_leadin_items_and_exception_together(self):
        records = upload_records(regulation_bytes())
        matches = [r for r in records if "주 15시간" in r.page_content]
        self.assertEqual(len(matches), 1)
        text = matches[0].page_content
        for fragment in ["제1장", "제2조", "지급 제외 사유", "제3조에 따른 예외"]:
            self.assertIn(fragment, text)

    def test_long_paragraph_is_not_silently_cut_or_overlapped(self):
        body = "① " + "대상자는 조건을 모두 충족해야 한다. " * 90 + "다만, 예외는 별도 심사한다."
        records = upload_records(regulation_bytes([("제1조(긴 조문)", [body])]))
        self.assertTrue(any(body in r.page_content for r in records))

    def test_duplicate_article_is_rejected_before_indexing(self):
        data = regulation_bytes([("제1조(가)", ["원문"]), ("제1조(나)", ["다른 원문"])])
        with self.assertRaises(ValueError):
            upload_records(data)

    def test_empty_article_is_rejected_before_indexing(self):
        with self.assertRaises(ValueError):
            upload_records(regulation_bytes([("제1조(빈 조문)", [])]))

    def test_legacy_heading_and_type_still_work(self):
        doc = Document()
        doc.add_paragraph("SKN34 학생 정책집")
        doc.add_heading("1-1 출결", level=1)
        doc.add_paragraph("정책 1-1 | 출결")
        doc.add_paragraph("지각 및 조퇴 합계 3회는 결석 1일로 환산한다.")
        buffer = BytesIO(); doc.save(buffer)
        records = build_records(load_docx(buffer.getvalue(), source_name="legacy.docx", cohort="cohort_34"))
        self.assertEqual(records[0].metadata["type"], "Attendance")
        self.assertIn("3회", records[0].page_content)

    def test_body_table_values_and_header_are_preserved(self):
        data = regulation_bytes([("제1조(출석률)", ["다음 표를 따른다."])])
        doc = Document(BytesIO(data))
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "구분"; table.cell(0, 1).text = "출석률"
        table.cell(1, 0).text = "장려금"; table.cell(1, 1).text = "80%"
        buffer = BytesIO(); doc.save(buffer)
        text = "\n".join(r.page_content for r in upload_records(buffer.getvalue()))
        for fragment in ["구분", "출석률", "장려금", "80%"]:
            self.assertIn(fragment, text)

    def test_upload_doc_ids_are_isolated_by_cohort_and_version(self):
        data = regulation_bytes()
        groups = [upload_records(data), upload_records(data, upload="b"),
                  upload_records(data, cohort="cohort_40")]
        ids = [{r.metadata["doc_id"] for r in group} for group in groups]
        self.assertFalse(ids[0] & ids[1])
        self.assertFalse(ids[0] & ids[2])

    def test_multiple_cohorts_can_be_built_together(self):
        data = regulation_bytes()
        sections = [s for cohort in ("cohort_34", "cohort_40")
                    for s in load_docx(data, source_name="rules.docx", cohort=cohort)]
        records = build_records(sections)
        self.assertEqual(len(records), 6)
        self.assertEqual(len({r.metadata['article_id'] for r in records}), 6)

    def test_explicit_reference_is_resolved_within_uploaded_document(self):
        records = upload_records(regulation_bytes())
        source = next(r for r in records if r.metadata['article_number'] == '2')
        target = next(r for r in records if r.metadata['article_number'] == '3')
        self.assertEqual(source.metadata['reference_vector_ids'], [target.vector_id])

    def test_external_and_unresolved_clause_references_do_not_link_locally(self):
        data = regulation_bytes([
            ("제1조(외부 근거)", ["근로기준법 제3조를 따른다.",
                                "「별도 운영 매뉴얼」 제3조에 따른다.",
                                "제99조 및 제3조제2항을 확인한다."]),
            ("제3조(내부 조문)", ["다른 내부 규정이다."]),
        ])
        source = upload_records(data)[0]
        self.assertEqual(source.metadata['reference_vector_ids'], [])
        self.assertTrue(source.metadata['unresolved_article_refs'])

    def test_inserted_article_number_does_not_resolve_to_base_article(self):
        records = upload_records(regulation_bytes([
            ("제1조(참조)", ["제3조의2를 따른다."]),
            ("제3조(기본)", ["기본 조문"]),
            ("제3조의2(추가)", ["추가 조문"]),
        ]))
        self.assertEqual(records[0].metadata['reference_vector_ids'], [records[2].vector_id])

    def test_draft_is_preserved_and_unmarked_document_is_not_approved(self):
        for title, expected in [("SKN34 운영규정(안)", "draft"), ("SKN34 운영규정", "unverified")]:
            records = upload_records(regulation_bytes(title=title))
            self.assertEqual({r.metadata['approval_status'] for r in records}, {expected})

    def test_numbered_cohort_input_is_normalized_for_search_filter(self):
        for value in ['34', '34기', 'cohort_34']:
            self.assertEqual(infer_cohort('', value), 'cohort_34')
        self.assertEqual(infer_cohort('', 'cohort-abc1234'), 'cohort-abc1234')

    def test_external_reference_list_does_not_link_second_article_locally(self):
        records = upload_records(regulation_bytes([
            ("제1조(근거)", ["근로기준법 제3조 및 제4조에 따른다."]),
            ("제3조(내부)", ["내부 규정"]), ("제4조(내부)", ["내부 규정"]),
        ]))
        self.assertEqual(records[0].metadata['reference_vector_ids'], [])

    def test_oversized_article_is_rejected_before_external_write(self):
        data = regulation_bytes([("제1조(과대 조문)", ["가" * 15000])])
        with self.assertRaises(ValueError):
            upload_records(data)

    def test_whole_article_preserves_condition_and_exception_across_long_units(self):
        paragraphs = ["① 다음 각 호를 충족해야 한다.", "1. " + "조건 " * 500,
                      "2. " + "조건 " * 500, "② 다만, 증빙 미제출자는 제외한다."]
        records = upload_records(regulation_bytes([("제1조(완결 조문)", paragraphs)]))
        self.assertEqual(len(records), 1)
        for paragraph in paragraphs:
            self.assertIn(paragraph.strip(), records[0].page_content)


if __name__ == "__main__":
    unittest.main()
