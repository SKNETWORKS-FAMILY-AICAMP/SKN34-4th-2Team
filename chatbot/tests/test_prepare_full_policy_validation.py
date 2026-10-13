import json
import tempfile
import unittest
from pathlib import Path
from docx import Document
from chatbot.evaluation.prepare_full_policy_validation import prepare

class PrepareFullPolicyTests(unittest.TestCase):
    def test_offline_copies_and_manifest_and_refuses_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'policy.docx'
            doc = Document()
            doc.add_heading('제1장 총칙', level=1)
            doc.add_heading('제1조 목적', level=2)
            doc.add_paragraph('학습 운영을 위한 규정입니다.')
            doc.save(source)
            original = source.read_bytes()
            output = root / 'output'
            manifest = prepare(policy_34=source, policy_40=source, output_dir=output, supplement=source)
            data = json.loads(manifest.read_text(encoding='utf-8'))
            self.assertEqual(data['external_calls'], 0)
            self.assertEqual(len(data['documents']), 3)
            for item in data['documents']:
                self.assertTrue(item['copy_hash_matches_source'])
                self.assertEqual((output / item['copy_path']).read_bytes(), original)
            self.assertEqual(data['documents'][2]['document_records']['status'], 'not_applicable')
            before = manifest.read_bytes()
            with self.assertRaises(FileExistsError):
                prepare(policy_34=source, policy_40=source, output_dir=output)
            self.assertEqual(manifest.read_bytes(), before)
            self.assertEqual(source.read_bytes(), original)

    def test_non_docx_input_rejected_without_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'policy.txt'
            source.write_text('test')
            with self.assertRaises(ValueError):
                prepare(policy_34=source, policy_40=source, output_dir=root / 'output')
            self.assertFalse((root / 'output').exists())
