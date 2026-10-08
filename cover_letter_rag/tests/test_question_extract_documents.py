"""Word · PowerPoint · 한글(hwpx) 양식 → 자기소개서 문항. 읽는 순서 · 형식 확인 · 압축 한도 · 글 없는 파일."""
import base64
import zipfile
from io import BytesIO

from app.question_extract import QuestionListOut, read_document
from tests.test_question_extract import FakeExtractor, post, q

DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
PPTX_TYPE = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
HWPX_TYPE = 'application/hwp+zip'

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
HP = 'http://www.hancom.co.kr/hwpml/2011/paragraph'


def data_url(kind, data):
    return f'data:{kind};base64,' + base64.b64encode(data).decode()


def make_docx(build):
    from docx import Document

    document = Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(files):
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, body in files.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def pptx_text(*paragraphs):
    body = ''.join(f'<a:p><a:r><a:t>{t}</a:t></a:r></a:p>' for t in paragraphs)
    return f'<p:sp><p:txBody>{body}</p:txBody></p:sp>'


def pptx_cell(text):
    return f'<a:tc><a:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></a:txBody></a:tc>'


def make_pptx(slides, order):
    """slides: 파일 이름 → 도형 XML. order: 발표 순서(파일 이름)"""
    rels = ''.join(f'<Relationship Id="rId{i}" Target="slides/{name}"/>' for i, name in enumerate(order, 1))
    ids = ''.join(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i, _ in enumerate(order, 1))
    files = {
        '[Content_Types].xml': '<Types/>',
        'ppt/presentation.xml': f'<p:presentation xmlns:p="{P}" xmlns:r="{R}"><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>',
        'ppt/_rels/presentation.xml.rels':
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>',
    }
    for name, body in slides.items():
        files[f'ppt/slides/{name}'] = f'<p:sld xmlns:p="{P}" xmlns:a="{A}"><p:cSld><p:spTree>{body}</p:spTree></p:cSld></p:sld>'
    return make_zip(files)


def hwpx_p(text, inner=''):
    return f'<hp:p><hp:run><hp:t>{text}</hp:t>{inner}</hp:run></hp:p>'


def hwpx_cell(text):
    return f'<hp:tc><hp:subList>{hwpx_p(text)}</hp:subList></hp:tc>'


def make_hwpx(sections, spine):
    names = list(sections)
    items = ''.join(f'<opf:item id="s{i}" href="Contents/{name}" media-type="application/xml"/>' for i, name in enumerate(names))
    refs = ''.join(f'<opf:itemref idref="s{names.index(name)}"/>' for name in spine)
    files = {
        'mimetype': 'application/hwp+zip',
        'Contents/content.hpf': (f'<opf:package xmlns:opf="http://www.idpf.org/2007/opf/">'
                                 f'<opf:manifest>{items}</opf:manifest><opf:spine>{refs}</opf:spine></opf:package>'),
    }
    for name, body in sections.items():
        files[f'Contents/{name}'] = f'<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="{HP}">{body}</hs:sec>'
    return make_zip(files)


def test_docx_paragraphs_and_tables_in_reading_order():
    def build(document):
        document.add_paragraph('자기소개서 양식')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '1. 지원 동기를 쓰시오.'
        table.cell(0, 1).text = '(500자 이내)'
        table.cell(1, 0).text = '2. 협업 경험을 쓰시오.'
        document.add_paragraph('3. 입사 후 포부를 쓰시오. (700자)')

    kind, text = read_document(DOCX_TYPE, make_docx(build))
    assert kind == 'docx'
    assert text.splitlines() == [
        '자기소개서 양식', '1. 지원 동기를 쓰시오. | (500자 이내)', '2. 협업 경험을 쓰시오.', '3. 입사 후 포부를 쓰시오. (700자)',
    ]


def test_pptx_follows_slide_order_not_file_names_and_reads_tables():
    table = ('<p:graphicFrame><a:graphic><a:graphicData><a:tbl>'
             f'<a:tr>{pptx_cell("문항")}{pptx_cell("글자 수")}</a:tr>'
             f'<a:tr>{pptx_cell("지원 동기")}{pptx_cell("500자")}</a:tr>'
             '</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    data = make_pptx({'slide1.xml': pptx_text('둘째 장'), 'slide2.xml': pptx_text('첫 장 제목', '안내') + table},
                     order=['slide2.xml', 'slide1.xml'])
    kind, text = read_document(PPTX_TYPE, data)
    assert kind == 'pptx'
    assert [line for line in text.splitlines() if line] == ['첫 장 제목', '안내', '문항 | 글자 수', '지원 동기 | 500자', '둘째 장']


def test_hwpx_follows_spine_and_reads_tables_inside_paragraphs():
    table = f'<hp:tbl><hp:tr>{hwpx_cell("1. 지원 동기")}{hwpx_cell("(600자)")}</hp:tr></hp:tbl>'
    data = make_hwpx({'section0.xml': hwpx_p('뒤 구역'), 'section1.xml': hwpx_p('자기소개서', inner=table)},
                     spine=['section1.xml', 'section0.xml'])
    kind, text = read_document(HWPX_TYPE, data)
    assert kind == 'hwpx'
    assert text.splitlines() == ['자기소개서', '1. 지원 동기 | (600자)', '뒤 구역']


def test_document_goes_through_the_same_text_path_and_is_grounded():
    def build(document):
        document.add_paragraph('1. 클라우드랩에 지원한 이유를 쓰시오. (1,000자)')

    extractor = FakeExtractor(QuestionListOut(questions=[q('클라우드랩에 지원한 이유를 쓰시오.', 1000), q('없는 문항을 쓰시오.')]))
    response = post(extractor, [data_url(DOCX_TYPE, make_docx(build))])
    assert response.status_code == 200, response.text
    assert response.json() == {'questions': [{'question': '클라우드랩에 지원한 이유를 쓰시오.', 'limit': 1000}],
                               'dropped': 1, 'read_text': True, 'source': 'docx'}
    assert extractor.images == []


def test_document_type_must_match_its_content():
    """종류 표기만 믿지 않는다 — Word 라고 왔는데 속이 PowerPoint 거나 ZIP 이 아니면 거절"""
    extractor = FakeExtractor(QuestionListOut())
    pptx = make_pptx({'slide1.xml': pptx_text('지원 동기')}, order=['slide1.xml'])
    for body in (pptx, b'PK\x03\x04 broken', b'not a zip'):
        response = post(extractor, [data_url(DOCX_TYPE, body)])
        assert response.status_code == 422 and response.json()['detail'] == 'document_unreadable', body[:10]
    assert extractor.texts == []


def test_document_without_text_is_not_an_empty_success():
    extractor = FakeExtractor(QuestionListOut())
    response = post(extractor, [data_url(DOCX_TYPE, make_docx(lambda document: None))])
    assert response.status_code == 422 and response.json()['detail'] == 'document_no_text'
    assert extractor.texts == []


def test_zip_bomb_is_refused_before_it_is_expanded():
    big = make_zip({'mimetype': 'application/hwp+zip', 'Contents/section0.xml': '<x/>' + ' ' * (25 * 1024 * 1024)})
    assert len(big) < 1024 * 1024
    response = post(FakeExtractor(QuestionListOut()), [data_url(HWPX_TYPE, big)])
    assert response.status_code == 422 and response.json()['detail'] == 'document_too_large'


def test_document_is_one_file_alone_and_hwp_is_not_a_document():
    extractor = FakeExtractor(QuestionListOut())
    docx = data_url(DOCX_TYPE, make_docx(lambda document: document.add_paragraph('지원 동기')))
    mixed = post(extractor, [docx, 'data:image/png;base64,AAAA'])
    assert mixed.status_code == 422 and mixed.json()['detail'] == 'document_with_other_files'
    assert post(extractor, [data_url('application/x-hwp', b'HWP Document File')]).status_code == 422
